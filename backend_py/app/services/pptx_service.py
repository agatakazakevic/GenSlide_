import os
from typing import Dict, Any
from io import BytesIO
import base64
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pptx.dml.color import RGBColor

UPLOAD_DIR = os.getenv('UPLOAD_DIR', './uploads')


def _slide_dimensions(prs: Presentation):
    width_in = prs.slide_width / 914400
    height_in = prs.slide_height / 914400
    return width_in, height_in


def _percent_to_inches(frame: Dict[str, float], slide_width: float, slide_height: float):
    return (
        Inches((frame['x'] / 100) * slide_width),
        Inches((frame['y'] / 100) * slide_height),
        Inches((frame['w'] / 100) * slide_width),
        Inches((frame['h'] / 100) * slide_height),
    )


def _add_chart(slide, chart_data: Dict[str, Any], frame, slide_width, slide_height):
    chart_type = (chart_data.get('type') or 'bar').lower()
    chart_map = {
        'bar': XL_CHART_TYPE.COLUMN_CLUSTERED,
        'line': XL_CHART_TYPE.LINE_MARKERS,
        'pie': XL_CHART_TYPE.PIE,
        'doughnut': XL_CHART_TYPE.DOUGHNUT,
    }

    chart_type_value = chart_map.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
    data = ChartData()
    data.categories = chart_data.get('labels') or []

    for dataset in chart_data.get('datasets') or []:
        data.add_series(dataset.get('label') or 'Series', dataset.get('data') or [])

    left, top, width, height = _percent_to_inches(frame, slide_width, slide_height)
    chart = slide.shapes.add_chart(chart_type_value, left, top, width, height, data).chart
    chart.has_legend = chart_data.get('showLegend', True)
    if chart_data.get('title'):
        chart.has_title = True
        chart.chart_title.text_frame.text = chart_data.get('title')


def create_editable_pptx(presentation_json: Dict[str, Any], images: Dict[str, Any]) -> Presentation:
    prs = Presentation()
    slide_width, slide_height = _slide_dimensions(prs)

    for slide_data in presentation_json.get('slides', []):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        if slide_data.get('backgroundColor'):
            fill = slide.background.fill
            fill.solid()
            hex_color = slide_data['backgroundColor'].replace('#', '')
            fill.fore_color.rgb = RGBColor.from_string(hex_color)

        if slide_data.get('title'):
            textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.7))
            textbox.text_frame.text = slide_data['title']

        if slide_data.get('subtitle'):
            textbox = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(9), Inches(0.5))
            textbox.text_frame.text = slide_data['subtitle']

        if slide_data.get('content'):
            y = 1.8
            for item in slide_data['content']:
                textbox = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(8.8), Inches(0.3))
                textbox.text_frame.text = f"• {item}"
                y += 0.35

        if slide_data.get('leftColumn', {}).get('type') == 'text':
            y = 1.6
            for item in slide_data['leftColumn']['items']:
                textbox = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(4.3), Inches(0.3))
                textbox.text_frame.text = f"• {item}"
                y += 0.35

        if slide_data.get('rightColumn', {}).get('type') == 'text':
            y = 1.6
            for item in slide_data['rightColumn']['items']:
                textbox = slide.shapes.add_textbox(Inches(5.1), Inches(y), Inches(4.3), Inches(0.3))
                textbox.text_frame.text = f"• {item}"
                y += 0.35

        if slide_data.get('leftColumn', {}).get('type') == 'chart':
            _add_chart(slide, slide_data['leftColumn']['chartData'], {"x": 6, "y": 28, "w": 43, "h": 60}, slide_width, slide_height)

        if slide_data.get('rightColumn', {}).get('type') == 'chart':
            _add_chart(slide, slide_data['rightColumn']['chartData'], {"x": 51, "y": 28, "w": 43, "h": 60}, slide_width, slide_height)

        if slide_data.get('chartData'):
            _add_chart(slide, slide_data['chartData'], {"x": 6, "y": 28, "w": 88, "h": 60}, slide_width, slide_height)

        if slide_data.get('table'):
            rows = [slide_data['table'].get('headers', [])] + slide_data['table'].get('rows', [])
            if rows:
                table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(1.6), Inches(9), Inches(3.6)).table
                for r_idx, row in enumerate(rows):
                    for c_idx, value in enumerate(row):
                        table.cell(r_idx, c_idx).text = str(value)

        if slide_data.get('image') and slide_data['image'].get('placeholderId'):
            placeholder_id = slide_data['image']['placeholderId']
            image_value = images.get(placeholder_id)
            if image_value and image_value.get('data'):
                frame = slide_data['image'].get('frame') or {"x": 0, "y": 0, "w": 100, "h": 100}
                left, top, width, height = _percent_to_inches(frame, slide_width, slide_height)
                data = image_value['data']
                if data.startswith('data:'):
                    base64_data = data.split(',', 1)[1]
                else:
                    base64_data = data
                image_bytes = BytesIO(base64.b64decode(base64_data))
                slide.shapes.add_picture(image_bytes, left, top, width, height)

    return prs


def save_pptx(prs: Presentation, filename: str) -> str:
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    filepath = os.path.join(UPLOAD_DIR, filename)
    prs.save(filepath)
    return filepath
